from pathlib import Path
from typing import Iterable
from uuid import uuid4
import os


os.environ.setdefault("USER_AGENT", "agentic-rag/1.0")

from langchain_community.document_loaders import (
    Docx2txtLoader,
    PyPDFLoader,
    TextLoader,
    WebBaseLoader,
)
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import load_workbook
from pptx import Presentation

from app.config import DATA_DIR


SUPPORTED_FILE_TYPES = {"pdf", "doc", "docx", "txt", "pptx", "xlsx"}


def _save_upload(file) -> tuple[Path, str]:
    ext = Path(file.name).suffix.lower().lstrip(".")
    target_dir = DATA_DIR / ext
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"{uuid4().hex}_{file.name}"
    target.write_bytes(file.read())
    return target, file.name


def _load_pptx(path: Path) -> list[Document]:
    presentation = Presentation(str(path))
    docs = []

    for index, slide in enumerate(presentation.slides, start=1):
        text = "\n".join(
            shape.text
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        )
        if text.strip():
            docs.append(
                Document(
                    page_content=text,
                    metadata={"source": path.name, "slide": index, "type": "pptx"},
                )
            )

    return docs


def _load_xlsx(path: Path) -> list[Document]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    docs = []

    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None]
            if values:
                rows.append(" | ".join(values))

        text = "\n".join(rows)
        if text.strip():
            docs.append(
                Document(
                    page_content=text,
                    metadata={"source": path.name, "sheet": sheet.title, "type": "xlsx"},
                )
            )

    return docs


def _load_file(path: Path, display_source: str | None = None) -> list[Document]:
    ext = path.suffix.lower().lstrip(".")
    source_name = display_source or path.name

    if ext == "pdf":
        docs = PyPDFLoader(str(path)).load()
    elif ext in {"doc", "docx"}:
        docs = Docx2txtLoader(str(path)).load()
    elif ext == "txt":
        docs = TextLoader(str(path), autodetect_encoding=True).load()
    elif ext == "pptx":
        docs = _load_pptx(path)
    elif ext == "xlsx":
        docs = _load_xlsx(path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    for doc in docs:
        doc.metadata["source"] = source_name
        doc.metadata.setdefault("type", ext)
        doc.metadata.setdefault("stored_file", path.name)

    return docs


def _load_urls(urls: Iterable[str]) -> list[Document]:
    docs = []
    for url in urls:
        cleaned = url.strip()
        if not cleaned:
            continue
        loaded = WebBaseLoader(cleaned).load()
        for doc in loaded:
            doc.metadata.setdefault("source", cleaned)
            doc.metadata.setdefault("type", "url")
        docs.extend(loaded)
    return docs


def load_documents(files=None, urls=None) -> list[Document]:
    documents = []

    for file in files or []:
        ext = Path(file.name).suffix.lower().lstrip(".")
        if ext not in SUPPORTED_FILE_TYPES:
            continue
        path, display_source = _save_upload(file)
        documents.extend(_load_file(path, display_source))

    documents.extend(_load_urls(urls or []))

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=700,
        chunk_overlap=120,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    return splitter.split_documents(documents)


def ingest_documents(files=None, urls=None, return_chunks: bool = False):
    from app.rag.vectordb import get_vector_db

    chunks = load_documents(files=files, urls=urls)
    if not chunks:
        return {"count": 0, "chunks": []} if return_chunks else 0

    db = get_vector_db()
    db.add_documents(chunks)
    db.persist()
    if return_chunks:
        return {"count": len(chunks), "chunks": chunks}
    return len(chunks)
